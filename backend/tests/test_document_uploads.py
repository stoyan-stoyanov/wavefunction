from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from conftest import ScriptedModel, agent_spec, tool_call_message
from fastapi import UploadFile
from langchain_core.messages import AIMessage, HumanMessage
from PIL import Image
from pydantic import Field

from wavefunction.agent_factory import build_agent
from wavefunction.desktop_files import build_desktop_file_tools
from wavefunction.spec_v2 import ToolBinding
from wavefunction.toolkit import ToolContext


def image_bytes(format="PNG", color="red"):
    output = BytesIO()
    Image.new("RGB", (16, 12), color).save(output, format=format)
    return output.getvalue()


async def upload(stack, owner, name, data):
    return await stack.files.save_upload(owner, UploadFile(BytesIO(data), filename=name))


@pytest.mark.parametrize("extension", ["png", "exe"])
async def test_invalid_uploads_leave_no_files_or_records(stack, extension):
    agent = stack.agent
    with pytest.raises(ValueError):
        await upload(stack, agent["id"], f"broken.{extension}", b"not a valid file")
    assert not stack.repository.list_project_files(agent["id"])
    assert not list((stack.files.project_dir(agent["id"]) / "sources").glob("*"))


class CapturingModel(ScriptedModel):
    seen: list = Field(default_factory=list)

    async def _agenerate(self, messages, stop=None, run_manager=None, **kwargs):
        self.seen.append(messages)
        return await super()._agenerate(messages, stop, run_manager, **kwargs)


def pixels(messages):
    return [
        block
        for message in messages
        if isinstance(message.content, list)
        for block in message.content
        if isinstance(block, dict) and block.get("type") == "image"
    ]


async def test_uploaded_image_reaches_the_model_in_real_agent_loops(stack):
    agent = stack.agent
    data = image_bytes()
    row = await upload(stack, agent["id"], "photo.png", data)
    context = ToolContext(
        agent["id"],
        stack.repository,
        stack.files,
        stack.secrets,
        reference_files=[row],
        coding_workspace=stack.services.coding_workspace,
    )
    assert row["media_type"].startswith("image/")
    assert (stack.files.project_dir(agent["id"]) / row["relative_path"]).read_bytes() == data
    assert "16 × 12" in stack.repository.get_project_file_texts(agent["id"], [row["id"]])[0]["extracted_text"]
    spec = agent_spec(ToolBinding(name="read_documents"))
    model = CapturingModel(
        script=[
            tool_call_message("read_documents", {"name": "photo.png"}, "read-image"),
            AIMessage(content="The reference shows red."),
        ]
    )
    stack.models.chat_model = lambda *_args, **_kwargs: model

    built = build_agent(spec, stack.models, context)
    result = await built.ainvoke({"messages": [HumanMessage(content="Inspect photo.png")]})
    assert result["messages"][-1].content == "The reference shows red."
    # Pixels are added to the provider request, not bloated into checkpoint history.
    assert not pixels(result["messages"])
    assert len(pixels(model.seen[-1])) == 1
    assert pixels(model.seen[-1])[0]["mime_type"] == "image/png"
    # A camera JPEG with an embedded preview opens as MPO; it is accepted as the JPEG it is.
    camera = BytesIO()
    Image.new("RGB", (64, 48), "red").save(
        camera, format="MPO", save_all=True, append_images=[Image.new("RGB", (16, 12), "blue")]
    )
    jpeg = await upload(stack, agent["id"], "camera.jpg", camera.getvalue())
    assert jpeg["media_type"] == "image/jpeg"
    # A phone photo larger than the preview limit still gets a thumbnail: JPEG decodes small.
    from wavefunction.document_images import PREVIEW_PIXELS, preview_png

    photo = BytesIO()
    Image.new("RGB", (5712, 4284), "blue").save(photo, format="JPEG", quality=30)
    assert PREVIEW_PIXELS < 5712 * 4284
    row = await upload(stack, agent["id"], "photo.jpg", photo.getvalue())
    stored = stack.files.project_dir(agent["id"]) / row["relative_path"]
    with Image.open(
        BytesIO(preview_png(stored, stored.stat().st_mtime_ns, stored.stat().st_size))
    ) as thumbnail:
        assert max(thumbnail.size) <= 320
    # Names that are not ASCII stay distinct for the user and the model.
    names = [(await upload(stack, agent["id"], name, b"text"))["name"] for name in ("報告.txt", "請求書.txt")]
    assert names == ["報告.txt", "請求書.txt"]
    assert (
        "64 × 48" in stack.repository.get_project_file_texts(agent["id"], [jpeg["id"]])[0]["extracted_text"]
    )


async def test_uploaded_slides_and_word_files_keep_text_in_tables_groups_charts_and_controls(stack):
    from docx import Document
    from docx.oxml import parse_xml
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "Quarterly numbers"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Region", "Revenue"
    table.cell(1, 0).text, table.cell(1, 1).text = "EMEA", "SYNTHETIC-42"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = "GROUPED-NOTE"
    numbers = CategoryChartData()
    numbers.categories = ["Q1", "Q2"]
    numbers.add_series("Sales", (1, 2))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(2), Inches(4), Inches(3), numbers
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "CHART-TITLE"
    slides = BytesIO()
    deck.save(slides)
    # A form's fields are content controls around ordinary paragraphs.
    document = Document()
    document.add_paragraph("Before the control")
    document.element.body.insert(
        1,
        parse_xml(
            '<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:sdtPr/>'
            "<w:sdtContent><w:p><w:r><w:t>CONTROL-TEXT</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        ),
    )
    word = BytesIO()
    document.save(word)

    rows = [
        await upload(stack, stack.agent["id"], "deck.pptx", slides.getvalue()),
        await upload(stack, stack.agent["id"], "form.docx", word.getvalue()),
    ]
    texts = stack.repository.get_project_file_texts(stack.agent["id"], [row["id"] for row in rows])
    by_name = {text["name"]: text["extracted_text"] for text in texts}
    assert by_name["deck.pptx"] == (
        "## Quarterly numbers\n- Region | Revenue\n- EMEA | SYNTHETIC-42\n- GROUPED-NOTE\n- CHART-TITLE\n"
    )
    assert by_name["form.docx"] == "Before the control\n\nCONTROL-TEXT\n"


async def test_desktop_transfers_preserve_files_and_only_read_captured_attachments(stack, tmp_path):
    agent = stack.agent
    reference = await upload(stack, agent["id"], "reference.txt", b"Attached to this task")
    await upload(stack, agent["id"], "other-chat.txt", b"Another chat's attachment")
    destination = tmp_path / "downloads"
    destination.mkdir()
    (destination / "report.txt").write_bytes(b"Existing user file")
    outside = tmp_path / "ungranted"
    outside.mkdir()
    stack.repository.set_setting("folder_grants", [{"path": str(destination), "writable": True}])
    session = object()
    computer = SimpleNamespace(
        running=lambda *_: session,
        session=AsyncMock(return_value=session),
        read_file=AsyncMock(return_value=b"Downloaded file"),
        add_file=AsyncMock(return_value="reference.txt"),
    )
    context = ToolContext(
        agent["id"],
        stack.repository,
        stack.files,
        stack.secrets,
        computer=computer,
        reference_files=[reference],
        coding_workspace=stack.services.coding_workspace,
    )
    tools = {tool.name: tool for tool in build_desktop_file_tools(context)}
    await tools["computer_add_file"].ainvoke({"source": "reference.txt"})
    computer.add_file.assert_awaited_once_with(session, "reference.txt", b"Attached to this task")
    with pytest.raises(ValueError, match="no document"):
        await tools["computer_add_file"].ainvoke({"source": "other-chat.txt"})
    await tools["computer_save_file"].ainvoke({"path": "report.txt", "destination": str(destination)})
    assert (destination / "report.txt").read_bytes() == b"Existing user file"
    assert (destination / "report (1).txt").read_bytes() == b"Downloaded file"
    with pytest.raises(ValueError, match="not allowed"):
        await tools["computer_save_file"].ainvoke({"path": "report.txt", "destination": str(outside)})
    assert not list(outside.iterdir())
